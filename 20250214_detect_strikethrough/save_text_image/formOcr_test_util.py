# -*- coding: utf-8 -*-
import numpy as np
import os
import json
import time
from queue import Queue
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Border, Side, Alignment
import cv2

exts = ['.jpg', '.png', '.tif', '.bmp']


def getImageListInDir (dir):
    file_list = os.listdir(dir)
    file_list.sort()
    file_list_img = Queue()
    for file in file_list:
        fname, fext = os.path.splitext(file)
        if fext.lower() in exts:
            file_list_img.put(os.path.join(dir, file))
    return file_list_img


def gen_sample_list(testCase, maxSample):
    process_sample_list = []

    if maxSample > 0:
        while len(process_sample_list) < maxSample:

            sample_data = []
            data = testCase["data"]
            for item in data:
                formOcrMode = item["formOcrMode"]
                formFile = item["formFile"]
                imageDirs = item["imageDir"]
                imageDirs.sort()
                img_list_queue = Queue()
                for dir in imageDirs:
                    img_list_queue_in_dir = getImageListInDir(dir)
                    while(img_list_queue_in_dir.qsize() > 0):
                        img_list_queue.put(img_list_queue_in_dir.get())
                sample_data.append({"formOcrMode": formOcrMode, "formFile": formFile, "imgs": img_list_queue})

            modified = True
            i = 0
            while modified and len(process_sample_list) < maxSample:
                modified = False
                subData = sample_data[i]
                img_queue = subData["imgs"]
                if img_queue.qsize() > 0:
                    process_sample_list.append([subData["formOcrMode"], subData["formFile"], img_queue.get()])
                    modified = True

                if i + 1 < len(sample_data):
                    i = i + 1
                else:
                    i = 0
    else:
        sample_data = []
        item = testCase["data"][0]
        formOcrMode = item["formOcrMode"]
        formFile = item["formFile"]
        imageDirs = item["imageDir"]
        imageDirs.sort()
        img_list_queue = Queue()
        for dir in imageDirs:
            img_list_queue_in_dir = getImageListInDir(dir)
            while(img_list_queue_in_dir.qsize() > 0):
                img_list_queue.put(img_list_queue_in_dir.get())
                process_sample_list.append([formOcrMode, formFile, img_list_queue.get()])

    return process_sample_list


def print_env (
    system_cpu_count,
    used_cpu_count,
    num_sample,
    dll_dir,
    model_dir,
    ocr_data_dir
    ):
    print('--------------------------------------------------------')
    print('InziFormOcr test')
    print('- processing info')
    print('  - system cpu count: ', system_cpu_count)
    print('  - used cpu count: ', used_cpu_count)
    print('  - sample count: ', num_sample)
    print('- env info')
    print('  - dll dir: ', dll_dir)
    print('  - model dir: ', model_dir)
    print('  - ocr data dir: ', ocr_data_dir)
    print('--------------------------------------------------------')


def init_ocr_result_file (ocr_result_file_path):
    if ocr_result_file_path:
        with open(ocr_result_file_path, "w", encoding='UTF8') as fp:
            pass


def log_ocr_result (ocr_result_file_path, result_list):
    if ocr_result_file_path:
        with open(ocr_result_file_path, "a", encoding='UTF8') as fp:
            for result in result_list:
                ec = result['ec']
                ocr_result = result['result_simple']
                if 0 == ec:
                    fp.write(ocr_result)
                else:
                    fp.write(f'image: {result[0]}\n')
                    fp.write(f'\terror: {ec}\n\n')


def init_time_log_file (time_log_file_path, system_cpu_count, used_cpu_count):
    if time_log_file_path:
        with open(time_log_file_path, "w", encoding='UTF8') as fp:
            fp.write(f'## Environment\n\n')
            fp.write(f'- system cpu count: {system_cpu_count}\n')
            fp.write(f'- used cpu count: {used_cpu_count}\n\n')
            fp.write(f'## Summary\n\n')
            title_data = "data count"
            title_avg = "avg"
            title_min = "min"
            title_max = "max"
            title_total = "total"
            title_test_case = "test case"
            fp.write(f'{title_data:>10}')
            fp.write(f'\t{title_avg:>8}\t{title_min:>8}\t{title_max:>8}\t{title_total:>8}')
            fp.write(f'\t{title_test_case}\n')
            fp.write('-'*80 + '\n')


def log_time (time_log_file_path, times, elepsed, data_count, test_name):
    if time_log_file_path:
        with open(time_log_file_path, "a", encoding='UTF8') as fp:
            fp.write(f'{data_count:10d}')
            fp.write(f'\t{elepsed/data_count: 8.3f}\t{times.min(): 8.3f}\t{times.max(): 8.3f}')
            fp.write(f'\t{elepsed: 8.3f}')
            fp.write(f'\t{test_name}\n')


def log_time_raw_data (time_log_file_path, elepsed, case_list, case_times):
    if time_log_file_path:
        with open(time_log_file_path, "a", encoding='UTF8') as fp:
            fp.write('-'*80 + '\n')
            # fp.write(f'{all_data_count:10d}')
            # fp.write(f'\t{np.mean(all_times): 8.3f}\t{min(all_times): 8.3f}\t{max(all_times): 8.3f}')
            # fp.write(f'\t{all_total_time: 8.3f}')
            # fp.write(f'\ttotal\n')
            fp.write(f'\n- Total testing time: {elepsed: 8.3f}\n')

            fp.write(f'\n## Raw data\n\n')
            for case, case_time in zip(case_list, case_times):
                fp.write(f'### Test case - {case["testName"]}\n\n')
                for img in case_time:
                    # print(img)
                    fp.write(f'{img[1]:8.3f}\t{img[0]}\t{img[2]}\n')
                fp.write('\n')


def save_result_json(result_dir, result_list):

    for result in result_list:
        image_file = result["img_file"]
        base = os.path.basename(image_file)
        json_file = os.path.join(result_dir, os.path.splitext(base)[0] + ".json")
        result_json = result["result_json"]

        # json 문자열 결과 저장
        with open(json_file, "w", encoding='UTF8') as fp:
            json.dump(result_json, fp, ensure_ascii=False, indent=2)


def save_text_image(result_dir, result_list):

    for result in result_list:
        image_file = result["img_file"]
        base = os.path.basename(image_file)
        text_file = os.path.join(result_dir, os.path.splitext(base)[0] + ".txt")
        img_text_file = os.path.join(result_dir, os.path.splitext(base)[0] + "_text.jpg")
        result_json = result["result_json"]

        # image_file 불러오기. 24비트 이미지로 불러오기
        img_array = np.fromfile(image_file, np.uint8)
        img = cv2.imdecode(img_array, cv2.IMREAD_COLOR)

        # text 정보 불러오기
        block_list = result_json["documents"][0]["blocks"]
        cnt_text = 0
        for block in block_list:
            for field in block["fields"]:
                text = field["text"]
                region = field["region"]

                cnt_text += 1
                crop_img_file_name = os.path.join(result_dir, os.path.splitext(base)[0] + "_" + str(cnt_text) + ".jpg")

                # 원본이미지에서 region 영역을 잘라서 저장
                img_text = img[region[1]:region[3], region[0]:region[2]]
                _, img_array = cv2.imencode('.jpg', img_text)
                img_array.tofile(crop_img_file_name)



def autofit_row_height(worksheet):
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value:
                # 1. Calculate the number of lines needed for the cell content
                lines = cell.value.count('\n') + 1

                # 2. Calculate the required height for the cell based on the default font size
                font_size = 11  # Change this value as needed
                cell_height = (font_size * 1.5) * lines

                # 3. Set the row height to fit the cell content
                if worksheet.row_dimensions[cell.row].height is None or worksheet.row_dimensions[cell.row].height < cell_height:
                    worksheet.row_dimensions[cell.row].height = cell_height


def autofit_column_width(worksheet):
    for column_cells in worksheet.columns:
        for cell in column_cells:
            if cell.value:
                # 4. Calculate the required width for the column based on the content length
                column_width = min(30, (len(str(cell.value)) * 1.2))

                # 5. Set the column width to fit the cell content
                if worksheet.column_dimensions[cell.column_letter].width is None or worksheet.column_dimensions[cell.column_letter].width < column_width:
                    worksheet.column_dimensions[cell.column_letter].width = column_width

def set_column_length_auto(ws):
    for col in ws.columns:
        new_column_length = max(len(str(cell.value)) for cell in col)
        new_column_length = min(30, new_column_length)
        new_column_letter = (get_column_letter(col[0].column))
        if new_column_length > 0:
            ws.column_dimensions[new_column_letter].width = new_column_length * 1.1

        # for cell in ws[new_column_letter]:
        #     cell.alignment = Alignment(horizontal='center')


def save_tables_to_excel(tables_data, output_filename):
    # Create a new Excel workbook
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Combined Table"

    # Define border style
    thin_border = Border(left=Side(style='thin'),
                         right=Side(style='thin'),
                         top=Side(style='thin'),
                         bottom=Side(style='thin'))

    current_row_offset = 0

    # Process each table
    for table_data in tables_data:
        # Find the maximum row used in the current table
        max_row = max(cell_info["end_row"] for cell_info in table_data["cells"])

        # Process each cell in the table data
        for cell_info in table_data["cells"]:
            start_row = cell_info["start_row"] + 1 + current_row_offset
            end_row = cell_info["end_row"] + 1 + current_row_offset
            start_col = cell_info["start_col"] + 1
            end_col = cell_info["end_col"] + 1
            value = cell_info["text"]

            # Set the value for the top-left cell
            ws.cell(row=start_row, column=start_col, value=value)
            ws.cell(row=start_row, column=start_col).alignment = openpyxl.styles.Alignment(wrap_text=True)

            # Merge cells if necessary
            if start_row != end_row or start_col != end_col:
                start_cell = f"{get_column_letter(start_col)}{start_row}"
                end_cell = f"{get_column_letter(end_col)}{end_row}"
                ws.merge_cells(f"{start_cell}:{end_cell}")

            # Apply border to all cells in the range
            for row in range(start_row, end_row + 1):
                for col in range(start_col, end_col + 1):
                    ws.cell(row=row, column=col).border = thin_border

        # Add offset for the next table
        current_row_offset += max_row + 2  # Leave one blank row between tables

    # autofit_column_width(ws)
    set_column_length_auto(ws)
    autofit_row_height(ws)

    # Save the workbook to the specified output file
    wb.save(output_filename)


def save_table_excel(result_dir, result_list):

    for result in result_list:
        image_file = result["img_file"]
        base = os.path.basename(image_file)
        xlsx_file = os.path.join(result_dir, os.path.splitext(base)[0] + ".xlsx")
        result_json = result["result_json"]

        tables = result_json["documents"][0]["tables"]

        # excel 문자열 결과 저장
        save_tables_to_excel(tables, xlsx_file)


def save_table_text(result_dir, result_list):

    for result in result_list:
        image_file = result["img_file"]
        base = os.path.basename(image_file)
        text_file = os.path.join(result_dir, os.path.splitext(base)[0] + ".txt")
        result_json = result["result_json"]

        tables = result_json["documents"][0]["tables"]

        with open(text_file, "w", encoding='UTF8') as fp:
            for table_data in tables:
                for cell_info in table_data["cells"]:
                    text = cell_info["text"]
                    if len(text) > 0:
                        fp.write(text + "\n")


def save_table_image(result_dir, result_list):

    for result in result_list:
        image_file = result["img_file"]
        base = os.path.basename(image_file)
        img_table_file = os.path.join(result_dir, os.path.splitext(base)[0] + "_table.jpg")
        result_json = result["result_json"]

        # image_file 불러오기. 24비트 이미지로 불러오기
        img_array = np.fromfile(image_file, np.uint8)
        img = cv2.imdecode(img_array, cv2.IMREAD_COLOR)

        tables = result_json["documents"][0]["tables"]
        for table in tables:
            table_region = table["region"]

            # margin 추가
            margin = 3
            table_region[0] = max(0, table_region[0] - margin)
            table_region[1] = max(0, table_region[1] - margin)
            table_region[2] = min(img.shape[1], table_region[2] + margin)
            table_region[3] = min(img.shape[0], table_region[3] + margin)

            # table_region을 img에 그리기, 색상은 초록색
            cv2.rectangle(img, (table_region[0], table_region[1]),
                          (table_region[2], table_region[3]),
                          (0, 255, 0), 2)

            # cell 그리기
            for cell in table["cells"]:
                cell_region = cell["region"]
                cv2.rectangle(img, (cell_region[0], cell_region[1]),
                              (cell_region[2], cell_region[3]),
                              (0, 0, 255), 3)

        # img 저장
        _, img_array = cv2.imencode('.jpg', img)
        img_array.tofile(img_table_file)
